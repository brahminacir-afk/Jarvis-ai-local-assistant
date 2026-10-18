import ctypes
import datetime as datetime_module
import json
import os
import re
import shutil
import socket
import subprocess
import threading
import uuid
import webbrowser
import zipfile
from urllib.parse import quote_plus
from xml.etree import ElementTree as ET

import requests

try:
    import winsound
except Exception:
    winsound = None

from ctypes import wintypes
from pathlib import Path

import psutil
from send2trash import send2trash

try:
    from PIL import ImageGrab
except Exception:  # pragma: no cover - optional fallback
    ImageGrab = None

try:
    from PIL import Image
except Exception:  # pragma: no cover - optional fallback
    Image = None

try:
    import winreg
except Exception:  # pragma: no cover - optional fallback
    winreg = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - optional fallback
    try:
        from PyPDF2 import PdfReader
    except Exception:  # pragma: no cover - optional fallback
        PdfReader = None

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - optional fallback
    load_workbook = None


BASE_DIR = Path(__file__).resolve().parent
with open(BASE_DIR / "config.json", "r", encoding="utf-8") as cfg:
    CONFIG = json.load(cfg)

HOME_DIR = Path.home()
DEFAULT_SAFE_DIR = HOME_DIR / "Documents"
DATA_ROOT = HOME_DIR / "Documents" / "Codex" / "SecureAssistantData"
NOTES_DIR = DATA_ROOT / "notes"
CAPTURE_DIR = DATA_ROOT / "captures"
SYSTEM_ACTIONS_DIR = DATA_ROOT / "system_actions"
ALARM_NOTIFY_SCRIPT = SYSTEM_ACTIONS_DIR / "alarm_notify.ps1"
ALARM_REGISTER_SCRIPT = SYSTEM_ACTIONS_DIR / "register_alarm_task.ps1"
OLLAMA_URL = CONFIG.get("ollama_url", "http://localhost:11434/api/generate")
ACTIVE_TIMERS = []
NOTIFICATION_CALLBACK = None

SAFE_DIR_MAP = {
    "desktop": HOME_DIR / "Desktop",
    "onedrive desktop": HOME_DIR / "OneDrive" / "Desktop",
    "onedrive\\desktop": HOME_DIR / "OneDrive" / "Desktop",
    "documents": HOME_DIR / "Documents",
    "downloads": HOME_DIR / "Downloads",
    "pictures": HOME_DIR / "Pictures",
    "project": BASE_DIR,
    "workspace": BASE_DIR,
}

VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".avi", ".webm"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".webp"}
TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".rst",
    ".py",
    ".cpp",
    ".c",
    ".cc",
    ".cxx",
    ".h",
    ".hh",
    ".hpp",
    ".inl",
    ".json",
    ".csv",
    ".js",
    ".ts",
    ".java",
    ".html",
    ".css",
    ".xml",
    ".yml",
    ".yaml",
    ".ini",
    ".cfg",
    ".toml",
    ".csv",
    ".tsv",
    ".ps1",
    ".bat",
    ".sh",
    ".tex",
    ".rtf",
}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm", ".ods"}
HIGH_RISK_INTENTS = {
    "delete_file",
    "move_file",
    "rename_file",
    "delete_folder",
    "move_folder",
    "rename_folder",
}
PROTECTED_ROOTS = [
    Path("C:/Windows"),
    Path("C:/Windows/System32"),
    Path("C:/Program Files"),
]


def set_notification_callback(callback):
    global NOTIFICATION_CALLBACK
    NOTIFICATION_CALLBACK = callback


def _notify(message):
    if NOTIFICATION_CALLBACK:
        try:
            NOTIFICATION_CALLBACK(message)
        except Exception:
            pass


def _fire_alert(kind, message, popup=False):
    prefix_map = {
        "alarm": "ALARM",
        "reminder": "REMINDER",
        "timer": "TIMER",
    }
    alert_prefix = prefix_map.get(kind, "ALERT")
    alert_text = f"{alert_prefix}: {message}"

    if winsound is not None:
        try:
            if kind == "alarm":
                winsound.MessageBeep(winsound.MB_ICONHAND)
                winsound.Beep(880, 250)
                winsound.Beep(988, 250)
            elif kind == "reminder":
                winsound.MessageBeep(winsound.MB_ICONEXCLAMATION)
            else:
                winsound.MessageBeep(winsound.MB_OK)
        except Exception:
            pass

    _notify(alert_text)

    if popup:
        try:
            ctypes.windll.user32.MessageBoxW(None, alert_text, alert_prefix, 0x00050030)
        except Exception:
            pass


def _get_parameter(intent_json, key, default=None):
    parameters = intent_json.get("parameters") or {}
    if key in parameters and parameters[key] not in (None, ""):
        return parameters[key]
    value = intent_json.get(key)
    if value not in (None, ""):
        return value
    return default


def _normalize_text(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _resolve_path_value(value, default_base=None):
    if not value:
        return None

    raw_text = str(value).strip().strip('"')
    path_key = raw_text.lower().replace('\\', ' ').replace('/', ' ').strip()
    if path_key in SAFE_DIR_MAP:
        return SAFE_DIR_MAP[path_key].resolve()

    parts = [part for part in re.split(r"[\\/]+", raw_text) if part]
    if parts:
        head_key = parts[0].lower()
        if head_key in SAFE_DIR_MAP:
            return SAFE_DIR_MAP[head_key].joinpath(*parts[1:]).resolve()
        if len(parts) >= 2:
            head_two_key = f"{parts[0].lower()} {parts[1].lower()}"
            if head_two_key in SAFE_DIR_MAP:
                return SAFE_DIR_MAP[head_two_key].joinpath(*parts[2:]).resolve()

    path = Path(os.path.expandvars(os.path.expanduser(raw_text)))
    if not path.is_absolute() and default_base:
        path = Path(default_base) / path
    return path.resolve()


def _is_within_allowed_root(path):
    try:
        resolved_path = Path(path).resolve()
    except Exception:
        return False

    for root in list(SAFE_DIR_MAP.values()) + [DEFAULT_SAFE_DIR]:
        try:
            resolved_path.relative_to(root.resolve())
            return True
        except Exception:
            continue
    return False


def _is_protected_path(path):
    try:
        resolved_path = Path(path).resolve()
    except Exception:
        return True

    for protected_root in PROTECTED_ROOTS:
        try:
            resolved_path.relative_to(protected_root.resolve())
            return True
        except Exception:
            continue
    return False


def _safe_target_path(raw_value, default_base=None):
    path = _resolve_path_value(raw_value, default_base=default_base)
    if path is None:
        return None
    if _is_protected_path(path):
        return None
    return path


def _resolve_alias_value(raw_value, alias_map):
    if raw_value is None:
        return raw_value
    return alias_map.get(str(raw_value), raw_value)


def _ensure_parent_directory(path):
    path.parent.mkdir(parents=True, exist_ok=True)


def _shell_open(target):
    try:
        result = ctypes.windll.shell32.ShellExecuteW(None, "open", str(target), None, None, 1)
        return result > 32
    except Exception:
        return False


def _resolve_app_target(target):
    normalized_target = _normalize_text(target)
    alias_map = {
        "notepad": "notepad.exe",
        "notepad++": "notepad++.exe",
        "opera": "opera.exe",
        "opera browser": "opera.exe",
        "google chrome": "chrome.exe",
        "chrome": "chrome.exe",
        "microsoft edge": "msedge.exe",
        "edge": "msedge.exe",
        "file explorer": "explorer.exe",
        "explorer": "explorer.exe",
        "calculator": "calc.exe",
        "paint": "mspaint.exe",
        "word": "winword.exe",
        "excel": "excel.exe",
        "powerpoint": "powerpnt.exe",
    }
    if normalized_target in alias_map:
        return alias_map[normalized_target]

    for alias, executable in alias_map.items():
        if alias in normalized_target:
            return executable

    raw_target = str(target or "").strip().strip('"')
    if not raw_target:
        return raw_target

    candidate_path = Path(os.path.expandvars(os.path.expanduser(raw_target)))
    if candidate_path.exists():
        return str(candidate_path)

    if candidate_path.suffix:
        return str(candidate_path)

    for candidate_name in (raw_target, f"{raw_target}.exe"):
        found = shutil.which(candidate_name)
        if found:
            return found

    registry_match = _find_app_path_registry_entry(raw_target)
    if registry_match:
        return registry_match

    shortcut_match = _find_start_menu_shortcut(raw_target)
    if shortcut_match:
        return shortcut_match

    executable_match = _find_installed_executable(raw_target)
    if executable_match:
        return executable_match

    found = shutil.which(raw_target)
    if found:
        return found

    return raw_target


def _find_app_path_registry_entry(target):
    if winreg is None:
        return None

    candidates = []
    base_name = Path(str(target)).stem.strip()
    if not base_name:
        return None

    if not base_name.lower().endswith(".exe"):
        candidates.append(f"{base_name}.exe")
    candidates.append(base_name)

    registry_roots = [
        (winreg.HKEY_CURRENT_USER, r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths"),
        (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths"),
        (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\WOW6432Node\Microsoft\Windows\CurrentVersion\App Paths"),
    ]

    for root, subkey in registry_roots:
        for candidate in candidates:
            try:
                with winreg.OpenKey(root, fr"{subkey}\{candidate}") as handle:
                    value, _ = winreg.QueryValueEx(handle, "")
                    if value:
                        return str(value)
            except Exception:
                continue
    return None


def _find_start_menu_shortcut(target):
    normalized_target = _normalize_text(target)
    if not normalized_target:
        return None

    start_menu_dirs = [
        os.path.join(os.environ.get("APPDATA", ""), r"Microsoft\Windows\Start Menu\Programs"),
        os.path.join(os.environ.get("ProgramData", ""), r"Microsoft\Windows\Start Menu\Programs"),
    ]

    for directory in start_menu_dirs:
        root = Path(directory)
        if not root.exists():
            continue
        for shortcut in root.rglob("*.lnk"):
            stem = _normalize_text(shortcut.stem)
            if normalized_target == stem or normalized_target in stem or stem in normalized_target:
                return str(shortcut)
    return None


def _find_installed_executable(target):
    normalized_target = _normalize_text(target)
    if not normalized_target:
        return None

    program_dirs = [
        os.environ.get("ProgramFiles"),
        os.environ.get("ProgramFiles(x86)"),
        os.environ.get("LOCALAPPDATA"),
    ]

    search_roots = []
    for base_dir in program_dirs:
        if not base_dir:
            continue
        search_roots.extend(
            [
                Path(base_dir),
                Path(base_dir) / "Programs",
            ]
        )

    seen_roots = []
    for root in search_roots:
        if not root.exists() or root in seen_roots:
            continue
        seen_roots.append(root)
        try:
            for candidate in root.rglob("*.exe"):
                stem = _normalize_text(candidate.stem)
                if normalized_target == stem or normalized_target in stem or stem in normalized_target:
                    return str(candidate)
        except Exception:
            continue
    return None


def _close_application(target):
    normalized_target = _normalize_text(target)
    if not normalized_target:
        return "No application name was provided."

    matches = []
    for process in psutil.process_iter(["name", "exe", "cmdline"]):
        try:
            process_name = _normalize_text(process.info.get("name", ""))
            process_exe = _normalize_text(Path(process.info.get("exe", "")).stem if process.info.get("exe") else "")
            command_line = " ".join(process.info.get("cmdline") or [])
            command_line = _normalize_text(command_line)
        except Exception:
            continue

        if (
            normalized_target == process_name
            or normalized_target == process_exe
            or normalized_target in process_name
            or normalized_target in process_exe
            or normalized_target in command_line
        ):
            matches.append(process)

    if not matches:
        return f"No running application matched '{target}'."

    closed_names = []
    for process in matches:
        try:
            process.terminate()
            closed_names.append(process.info.get("name") or str(process.pid))
        except Exception:
            continue

    gone, alive = psutil.wait_procs(matches, timeout=3)
    for process in alive:
        try:
            process.kill()
        except Exception:
            pass

    if closed_names:
        return f"Closed: {', '.join(sorted(set(closed_names)))}"
    return f"Attempted to close '{target}'."


def _open_folder(path):
    if path is None:
        return "No folder path was provided."
    if not Path(path).exists():
        return f"Folder not found: {path}"
    _shell_open(path)
    return f"Opened folder: {path}"


def _open_file(file_path):
    if file_path is None:
        return "No file path was provided."
    if not Path(file_path).exists():
        return f"File not found: {file_path}"
    _shell_open(file_path)
    return f"Opened file: {file_path}"


def _create_folder(folder_path):
    if folder_path is None:
        return "No folder path was provided."
    if _is_protected_path(folder_path):
        return f"Blocked protected path: {folder_path}"
    if not _is_within_allowed_root(folder_path):
        return f"Blocked unsafe destination outside allowed folders: {folder_path}"

    try:
        Path(folder_path).mkdir(parents=True, exist_ok=True)
        return f"Created folder: {folder_path}"
    except Exception as error:
        return f"Folder creation failed: {error}"


def _list_folder(folder_path, limit=20):
    if folder_path is None:
        return "No folder path was provided."
    path = Path(folder_path)
    if not path.exists():
        return f"Folder not found: {folder_path}"
    if not path.is_dir():
        return f"Not a folder: {folder_path}"

    items = sorted(path.iterdir(), key=lambda item: (not item.is_dir(), item.name.lower()))
    if not items:
        return f"{path} is empty."

    lines = []
    for item in items[:limit]:
        prefix = "[DIR]" if item.is_dir() else "[FILE]"
        lines.append(f"{prefix} {item.name}")

    remaining = len(items) - limit
    summary = f"{path} contains {len(items)} item(s):"
    if remaining > 0:
        lines.append(f"... and {remaining} more")
    return summary + "\n" + "\n".join(lines)


def _rename_folder(source_path, new_name):
    if source_path is None or not new_name:
        return "Source path or new folder name missing."
    if _is_protected_path(source_path):
        return f"Blocked protected path: {source_path}"

    source = Path(source_path)
    if not source.exists():
        return f"Folder not found: {source_path}"
    if not source.is_dir():
        return f"Not a folder: {source_path}"

    new_name = Path(str(new_name)).name
    destination = source.with_name(new_name)
    if not _is_within_allowed_root(source) or not _is_within_allowed_root(destination):
        return "Blocked unsafe rename outside allowed folders."

    try:
        source.rename(destination)
        return f"Renamed folder to: {destination}"
    except Exception as error:
        return f"Folder rename failed: {error}"


def _move_or_copy_folder(source_path, destination_path, operation):
    if source_path is None or destination_path is None:
        return "Source or destination path missing."
    if _is_protected_path(source_path) or _is_protected_path(destination_path):
        return "Blocked protected path in folder transfer request."

    source = Path(source_path)
    destination = Path(destination_path)
    if not source.exists():
        return f"Folder not found: {source_path}"
    if not source.is_dir():
        return f"Not a folder: {source_path}"
    if destination.exists() and destination.is_dir():
        destination = destination / source.name
    if not _is_within_allowed_root(source) or not _is_within_allowed_root(destination):
        return "Blocked unsafe folder transfer outside allowed folders."

    try:
        if operation == "move":
            shutil.move(str(source), str(destination))
            return f"Moved folder '{source}' to '{destination}'"
        shutil.copytree(str(source), str(destination), dirs_exist_ok=True)
        return f"Copied folder '{source}' to '{destination}'"
    except Exception as error:
        return f"Folder transfer failed: {error}"


def _delete_folder(folder_path):
    if folder_path is None:
        return "No folder path was provided."
    if _is_protected_path(folder_path):
        return f"Blocked protected path: {folder_path}"

    path = Path(folder_path)
    if not path.exists():
        return f"Folder not found: {folder_path}"
    if not path.is_dir():
        return f"Not a folder: {folder_path}"
    if not _is_within_allowed_root(path):
        return "Blocked unsafe folder deletion outside allowed folders."

    try:
        send2trash(str(path))
        return f"Moved folder to Recycle Bin: {path}"
    except Exception as error:
        return f"Folder delete failed: {error}"


def _search_files(folder_path, query, limit=20):
    if folder_path is None:
        return "No folder path was provided."
    path = Path(folder_path)
    if not path.exists():
        return f"Folder not found: {folder_path}"
    if not path.is_dir():
        return f"Not a folder: {folder_path}"

    query_text = _normalize_text(query)
    if not query_text:
        return "No search query was provided."

    matches = []
    for candidate in path.rglob("*"):
        if not candidate.is_file():
            continue
        candidate_text = _normalize_text(candidate.name)
        if query_text in candidate_text or query_text in _normalize_text(str(candidate)):
            matches.append(candidate)
        if len(matches) >= limit:
            break

    if not matches:
        return f"No matching files were found in {path}."

    lines = [str(item) for item in matches]
    return f"Found {len(matches)} matching file(s):\n" + "\n".join(lines)


def _create_file(file_path, content=""):
    if file_path is None:
        return "No file path was provided."
    if _is_protected_path(file_path):
        return f"Blocked protected path: {file_path}"
    if not _is_within_allowed_root(file_path):
        return f"Blocked unsafe destination outside allowed folders: {file_path}"

    try:
        _ensure_parent_directory(file_path)
        Path(file_path).write_text(content or "", encoding="utf-8")
        return f"Created file: {file_path}"
    except Exception as error:
        return f"File creation failed: {error}"


def _move_or_copy_file(source_path, destination_path, operation):
    if source_path is None or destination_path is None:
        return "Source or destination path missing."
    if _is_protected_path(source_path) or _is_protected_path(destination_path):
        return "Blocked protected path in file transfer request."
    if not Path(source_path).exists():
        return f"Source file not found: {source_path}"
    if not _is_within_allowed_root(source_path) or not _is_within_allowed_root(destination_path):
        return "Blocked unsafe file transfer outside allowed folders."

    try:
        _ensure_parent_directory(destination_path)
        if operation == "move":
            shutil.move(str(source_path), str(destination_path))
            return f"Moved '{source_path}' to '{destination_path}'"
        shutil.copy2(str(source_path), str(destination_path))
        return f"Copied '{source_path}' to '{destination_path}'"
    except Exception as error:
        return f"File transfer failed: {error}"


def _rename_file(source_path, new_name):
    if source_path is None or not new_name:
        return "Source path or new name missing."
    if _is_protected_path(source_path):
        return f"Blocked protected path: {source_path}"
    if not Path(source_path).exists():
        return f"Source file not found: {source_path}"

    new_name = Path(str(new_name)).name
    destination_path = Path(source_path).with_name(new_name)
    if not _is_within_allowed_root(source_path) or not _is_within_allowed_root(destination_path):
        return "Blocked unsafe rename outside allowed folders."

    try:
        Path(source_path).rename(destination_path)
        return f"Renamed '{source_path}' to '{destination_path.name}'"
    except Exception as error:
        return f"Rename failed: {error}"


def _find_latest_file(folder_path, file_kind=None):
    if folder_path is None:
        return "No folder path was provided."
    if not Path(folder_path).exists():
        return f"Folder not found: {folder_path}"

    candidates = []
    for candidate in Path(folder_path).rglob("*"):
        if not candidate.is_file():
            continue
        suffix = candidate.suffix.lower()
        if file_kind == "video" and suffix not in VIDEO_EXTENSIONS:
            continue
        if file_kind == "image" and suffix not in IMAGE_EXTENSIONS:
            continue
        candidates.append(candidate)

    if not candidates:
        return f"No matching files were found in {folder_path}."

    latest_file = max(candidates, key=lambda item: item.stat().st_mtime)
    _shell_open(latest_file)
    return f"Latest file found: {latest_file}"


def _read_document_text(file_path, max_chars=2500):
    path = Path(file_path)
    suffix = path.suffix.lower()

    try:
        if suffix == ".docx":
            return _read_docx_text(path, max_chars=max_chars)

        if suffix == ".odt":
            return _read_odt_text(path, max_chars=max_chars)

        if suffix in SPREADSHEET_EXTENSIONS:
            return _read_spreadsheet_text(path, max_chars=max_chars)

        if suffix == ".pdf" and PdfReader is not None:
            reader = PdfReader(str(path))
            fragments = []
            for page in reader.pages[:5]:
                try:
                    fragments.append(page.extract_text() or "")
                except Exception:
                    continue
            text = "\n".join(fragments).strip()
            return text[:max_chars] if text else None

        if suffix not in TEXT_EXTENSIONS:
            return None

        text = path.read_text(encoding="utf-8", errors="ignore").strip()
        if not text:
            return None
        return text[:max_chars]
    except Exception:
        return None


def _read_spreadsheet_text(path, max_chars=2500):
    suffix = path.suffix.lower()

    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"} and load_workbook is not None:
        try:
            workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
            fragments = []
            for sheet in workbook.worksheets[:5]:
                fragments.append(f"Sheet: {sheet.title}\n")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value not in (None, "")]
                    if values:
                        fragments.append(" | ".join(values) + "\n")
                fragments.append("\n")
            text = "".join(fragments).strip()
            return text[:max_chars] if text else None
        except Exception:
            return None

    if suffix in {".csv", ".tsv"}:
        try:
            delimiter = "\t" if suffix == ".tsv" else ","
            lines = path.read_text(encoding="utf-8", errors="ignore").splitlines()
            preview = []
            for line in lines[:120]:
                parts = [part.strip() for part in line.split(delimiter)]
                values = [part for part in parts if part]
                if values:
                    preview.append(" | ".join(values))
            text = "\n".join(preview).strip()
            return text[:max_chars] if text else None
        except Exception:
            return None

    return None


def _read_docx_text(path, max_chars=2500):
    try:
        with zipfile.ZipFile(path) as archive:
            with archive.open("word/document.xml") as handle:
                xml_data = handle.read()
    except Exception:
        return None

    try:
        root = ET.fromstring(xml_data)
    except Exception:
        return None

    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    fragments = []
    for node in root.iter():
        if node.tag == f"{namespace}t" and node.text:
            fragments.append(node.text)
        elif node.tag == f"{namespace}tab":
            fragments.append("\t")
        elif node.tag == f"{namespace}br":
            fragments.append("\n")
        elif node.tag == f"{namespace}p":
            fragments.append("\n")

    text = re.sub(r"\n{3,}", "\n\n", "".join(fragments)).strip()
    return text[:max_chars] if text else None


def _read_odt_text(path, max_chars=2500):
    try:
        with zipfile.ZipFile(path) as archive:
            with archive.open("content.xml") as handle:
                xml_data = handle.read()
    except Exception:
        return None

    try:
        root = ET.fromstring(xml_data)
    except Exception:
        return None

    fragments = []
    for node in root.iter():
        tag = node.tag.split("}", 1)[-1] if "}" in node.tag else node.tag
        if tag == "p" and node.text:
            fragments.append(node.text)
            fragments.append("\n")
        elif tag == "span" and node.text:
            fragments.append(node.text)
        elif tag == "tab":
            fragments.append("\t")
        elif tag == "s":
            fragments.append(" ")

    text = re.sub(r"\n{3,}", "\n\n", "".join(fragments)).strip()
    return text[:max_chars] if text else None


def _collect_folder_samples(folder_path, max_files=12, max_chars_per_file=2500, max_total_chars=18000):
    path = Path(folder_path)
    candidates = []
    for candidate in path.rglob("*"):
        if candidate.is_file():
            candidates.append(candidate)

    candidates.sort(key=lambda item: item.relative_to(path).as_posix().lower())

    samples = []
    total_chars = 0
    for candidate in candidates:
        if len(samples) >= max_files or total_chars >= max_total_chars:
            break

        content = _read_document_text(candidate, max_chars=max_chars_per_file)
        if not content:
            continue

        relative_name = candidate.relative_to(path).as_posix()
        samples.append(
            {
                "path": relative_name,
                "suffix": candidate.suffix.lower().lstrip("."),
                "content": content,
            }
        )
        total_chars += len(content)

    return samples


def _build_folder_study_prompt(folder_path, request_text, analysis_mode, samples):
    sample_blocks = []
    for sample in samples:
        sample_blocks.append(
            f"File: {sample['path']}\nType: {sample['suffix'] or 'unknown'}\nContent:\n{sample['content']}"
        )

    joined_samples = "\n\n---\n\n".join(sample_blocks) if sample_blocks else "No readable text files were found."
    return f"""
You are a helpful study assistant for local files.
Read the sampled course material and answer the user's request clearly.

User request: {request_text}
Requested style: {analysis_mode}
Folder: {folder_path}

Rules:
- Use simple language.
- If the material looks technical, explain it like a beginner-friendly tutor.
- If the user asked for exercises, provide a short set of practice exercises and keep them practical.
- If the content is incomplete, say what you could and could not inspect.
- Do not invent details that are not supported by the files.
- Prefer organized headings and short bullets.

Sampled content:
{joined_samples}
""".strip()


def _collect_upload_samples(upload_paths, max_files=18, max_chars_per_file=3500, max_total_chars=26000):
    samples = []
    total_chars = 0
    skipped = []

    for upload_path in upload_paths:
        path = Path(upload_path)
        if not path.exists():
            skipped.append(f"{path} (not found)")
            continue

        if path.is_file():
            candidates = [path]
            root = path.parent
        elif path.is_dir():
            candidates = sorted(
                (candidate for candidate in path.rglob("*") if candidate.is_file()),
                key=lambda item: item.relative_to(path).as_posix().lower(),
            )
            root = path
        else:
            skipped.append(f"{path} (not a regular file or folder)")
            continue

        for candidate in candidates:
            if len(samples) >= max_files or total_chars >= max_total_chars:
                skipped.append(f"{path} (sample limit reached)")
                break

            content = _read_document_text(candidate, max_chars=max_chars_per_file)
            if not content:
                skipped.append(f"{candidate} (unreadable or unsupported)")
                continue

            try:
                display_path = candidate.relative_to(root).as_posix() if path.is_dir() else candidate.name
            except Exception:
                display_path = str(candidate)

            samples.append(
                {
                    "path": str(path),
                    "display_path": display_path,
                    "suffix": candidate.suffix.lower().lstrip("."),
                    "content": content,
                }
            )
            total_chars += len(content)

    return samples, skipped


def _build_upload_prompt(upload_paths, request_text, samples, skipped):
    sample_blocks = []
    for sample in samples:
        sample_blocks.append(
            "Source: {source}\nFile: {file_name}\nType: {file_type}\nContent:\n{content}".format(
                source=sample["path"],
                file_name=sample["display_path"],
                file_type=sample["suffix"] or "unknown",
                content=sample["content"],
            )
        )

    joined_samples = "\n\n---\n\n".join(sample_blocks)
    skipped_text = "\n".join(f"- {item}" for item in skipped[:20]) if skipped else "None"

    return f"""
You are Jarvis, a local assistant answering questions about uploaded local file content.
The user asked: {request_text}

Uploaded path(s):
{chr(10).join(f"- {path}" for path in upload_paths)}

Rules:
- Answer only from the provided file content.
- If the user asks for code explanation, explain the purpose, important functions, flow, and any issues you notice.
- If the user asks for a summary, keep it structured and practical.
- Mention when content was sampled or when files could not be read.
- Do not invent files, APIs, or behavior that are not visible in the content.

Skipped or unreadable items:
{skipped_text}

Readable content:
{joined_samples}
""".strip()


def _generate_upload_response(upload_paths, request_text):
    samples, skipped = _collect_upload_samples(upload_paths)
    if not samples:
        if skipped:
            return "I could not read any supported text from the upload target(s):\n" + "\n".join(f"- {item}" for item in skipped[:20])
        return "I could not find readable content in the upload target(s)."

    prompt = _build_upload_prompt(upload_paths, request_text, samples, skipped)
    payload = {
        "model": CONFIG["model_name"],
        "prompt": prompt,
        "stream": False,
    }

    try:
        response = requests.post(OLLAMA_URL, json=payload, timeout=120)
        response.raise_for_status()
        body = response.json()
        answer = str(body.get("response", "")).strip()
        if answer:
            return answer
        return "I read the upload target(s), but the local model returned an empty response."
    except Exception as error:
        sample_lines = [f"- {sample['display_path']}" for sample in samples]
        return (
            f"I read {len(samples)} file sample(s), but Ollama was unavailable: {error}\n"
            f"Files inspected:\n" + "\n".join(sample_lines)
        )


def _resolve_upload_paths(intent_json, alias_map):
    raw_paths = _get_parameter(intent_json, "paths")
    if not raw_paths:
        raw_paths = [
            _get_parameter(intent_json, "path")
            or _get_parameter(intent_json, "folder")
            or _get_parameter(intent_json, "target")
        ]
    elif isinstance(raw_paths, str):
        raw_paths = [raw_paths]

    resolved_paths = []
    for raw_path in raw_paths or []:
        if not raw_path:
            continue
        alias_value = alias_map.get(str(raw_path))
        if alias_value:
            raw_path = alias_value
        resolved_path = _safe_target_path(raw_path, default_base=DEFAULT_SAFE_DIR)
        if not resolved_path:
            continue
        if not _is_within_allowed_root(resolved_path):
            continue
        resolved_paths.append(resolved_path)

    return resolved_paths


def _generate_folder_study_response(folder_path, request_text, analysis_mode):
    samples = _collect_folder_samples(folder_path)
    if not samples:
        return f"I could not find readable study files in {folder_path}."

    prompt = _build_folder_study_prompt(folder_path, request_text, analysis_mode, samples)
    payload = {
        "model": CONFIG["model_name"],
        "prompt": prompt,
        "stream": False,
    }

    try:
        response = requests.post(OLLAMA_URL, json=payload, timeout=120)
        response.raise_for_status()
        body = response.json()
        answer = str(body.get("response", "")).strip()
        if answer:
            return answer
        return "I read the files, but the local model returned an empty response."
    except Exception as error:
        summary_lines = [f"- {sample['path']}" for sample in samples]
        return (
            f"I read {len(samples)} file(s) from {folder_path}, but I could not generate the study response "
            f"because the local model was unavailable: {error}\n"
            f"Files inspected:\n" + "\n".join(summary_lines)
        )


def _resolve_analysis_folder(intent_json, alias_map):
    folder_value = _get_parameter(intent_json, "folder") or _get_parameter(intent_json, "path") or _get_parameter(intent_json, "target")
    if folder_value:
        alias_value = alias_map.get(str(folder_value))
        if alias_value:
            folder_value = alias_value
        folder_path = _safe_target_path(folder_value, default_base=DEFAULT_SAFE_DIR)
        if folder_path and folder_path.exists() and folder_path.is_dir() and _is_within_allowed_root(folder_path):
            return folder_path
    return None


def _save_note(note_text):
    if not note_text:
        return "No note text was provided."
    NOTES_DIR.mkdir(parents=True, exist_ok=True)
    note_file = NOTES_DIR / f"notes_{datetime_module.datetime.now().date().isoformat()}.txt"
    timestamp = datetime_module.datetime.now().strftime("%H:%M:%S")
    with open(note_file, "a", encoding="utf-8") as handle:
        handle.write(f"[{timestamp}] {note_text}\n")
    return f"Saved note to {note_file.name}"


def _read_note():
    note_file = NOTES_DIR / f"notes_{datetime_module.datetime.now().date().isoformat()}.txt"
    if not note_file.exists():
        return "No notes were found for today."
    return note_file.read_text(encoding="utf-8").strip() or "No notes were found for today."


def _clear_note():
    note_file = NOTES_DIR / f"notes_{datetime_module.datetime.now().date().isoformat()}.txt"
    note_file.parent.mkdir(parents=True, exist_ok=True)
    note_file.write_text("", encoding="utf-8")
    return "Cleared today's notes."


def _normalize_window_name(value):
    normalized = _normalize_text(value)
    normalized = normalized.replace("window", " ").replace("application", " ").replace("app", " ")
    return re.sub(r"\s+", " ", normalized).strip()


def _find_window_handle(target_text=None, window_title=None):
    user32 = ctypes.windll.user32
    target_value = _normalize_window_name(window_title or target_text)
    if not target_value:
        return None

    matches = []

    @ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p)
    def enum_windows_proc(hwnd, lparam):
        if not user32.IsWindowVisible(hwnd):
            return True

        length = user32.GetWindowTextLengthW(hwnd)
        if length <= 0:
            return True

        buffer = ctypes.create_unicode_buffer(length + 1)
        user32.GetWindowTextW(hwnd, buffer, length + 1)
        title = buffer.value.strip()
        if not title:
            return True

        normalized_title = _normalize_window_name(title)
        if target_value == normalized_title or target_value in normalized_title or normalized_title in target_value:
            matches.append(hwnd)
        return True

    user32.EnumWindows(enum_windows_proc, 0)
    if matches:
        return matches[0]
    return None


def _capture_window_to_path(hwnd, output_path):
    if hwnd is None or ImageGrab is None:
        return False

    user32 = ctypes.windll.user32
    try:
        rect = wintypes.RECT()
        if not user32.GetWindowRect(hwnd, ctypes.byref(rect)):
            return False

        width = rect.right - rect.left
        height = rect.bottom - rect.top
        if width <= 0 or height <= 0:
            return False

        image = ImageGrab.grab(bbox=(rect.left, rect.top, rect.right, rect.bottom), all_screens=True)
        if image is None:
            return False
        image.save(output_path)
        return True
    except Exception:
        return False


def _capture_screen(scope="full", filename_prefix="capture", target=None, window_title=None):
    if ImageGrab is None:
        return "Screenshot support is unavailable because Pillow is not installed."

    CAPTURE_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime_module.datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = CAPTURE_DIR / f"{filename_prefix}_{timestamp}.png"

    try:
        normalized_scope = _normalize_text(scope)
        if normalized_scope == "window":
            hwnd = _find_window_handle(target_text=target, window_title=window_title)
            if hwnd is None:
                return f"Could not find a window matching '{window_title or target}'."
            if _capture_window_to_path(hwnd, output_path):
                return f"Saved window capture to {output_path}"
            return f"Capture failed for window: {window_title or target}"

        image = ImageGrab.grab(all_screens=True)
        if normalized_scope == "active":
            user32 = ctypes.windll.user32
            hwnd = user32.GetForegroundWindow()
            rect = wintypes.RECT()
            if user32.GetWindowRect(hwnd, ctypes.byref(rect)):
                left, top, right, bottom = rect.left, rect.top, rect.right, rect.bottom
                image = image.crop((left, top, right, bottom))
        image.save(output_path)
        return f"Saved capture to {output_path}"
    except Exception as error:
        return f"Capture failed: {error}"


def _schedule_timer(message, delay_seconds):
    if delay_seconds is None:
        return "No timer duration was provided."
    delay_seconds = max(0, float(delay_seconds))

    def _fire():
        _fire_alert("timer", message)

    timer = threading.Timer(delay_seconds, _fire)
    timer.daemon = True
    timer.start()
    ACTIVE_TIMERS.append(timer)
    return f"Timer scheduled for {int(delay_seconds)} seconds."


def _ensure_system_alarm_scripts():
    SYSTEM_ACTIONS_DIR.mkdir(parents=True, exist_ok=True)

    notify_script = SYSTEM_ACTIONS_DIR / "alarm_notify.ps1"
    register_script = SYSTEM_ACTIONS_DIR / "register_alarm_task.ps1"

    notify_script_content = """param(
    [Parameter(Mandatory = $true)]
    [ValidateSet('alarm', 'reminder')]
    [string]$Kind,

    [Parameter(Mandatory = $true)]
    [string]$Title,

    [Parameter(Mandatory = $true)]
    [string]$Message
)

try {
    if ($Kind -eq 'alarm') {
        [Windows.UI.Notifications.ToastNotificationManager, Windows.UI.Notifications, ContentType=WindowsRuntime] | Out-Null
        [Windows.Data.Xml.Dom.XmlDocument, Windows.Data.Xml.Dom, ContentType=WindowsRuntime] | Out-Null
        $safeTitle = [System.Security.SecurityElement]::Escape($Title)
        $safeMessage = [System.Security.SecurityElement]::Escape($Message)
        $xml = [Windows.Data.Xml.Dom.XmlDocument]::new()
        $xml.LoadXml('<toast scenario="alarm"><visual><binding template="ToastGeneric"><text>' + $safeTitle + '</text><text>' + $safeMessage + '</text></binding></visual><audio src="ms-winsoundevent:Notification.Looping.Alarm" loop="true"/><actions><action content="Snooze" arguments="snooze" activationType="system"/><action content="Dismiss" arguments="dismiss" activationType="system"/></actions></toast>')
        $toast = [Windows.UI.Notifications.ToastNotification]::new($xml)
        $appId = '{1AC14E77-02E7-4E5D-B744-2EB1AE5198B7}\\WindowsPowerShell\\v1.0\\powershell.exe'
        [Windows.UI.Notifications.ToastNotificationManager]::CreateToastNotifier($appId).Show($toast)
    } else {
        Add-Type -AssemblyName System.Windows.Forms
        [System.Media.SystemSounds]::Asterisk.Play()
        [System.Windows.Forms.MessageBox]::Show($Message, $Title, [System.Windows.Forms.MessageBoxButtons]::OK, [System.Windows.Forms.MessageBoxIcon]::Information) | Out-Null
    }
} catch {
}
"""

    register_script_content = """param(
    [Parameter(Mandatory = $true)]
    [string]$TaskName,

    [Parameter(Mandatory = $true)]
    [string]$RunAtIso,

    [Parameter(Mandatory = $true)]
    [ValidateSet('alarm', 'reminder')]
    [string]$Kind,

    [Parameter(Mandatory = $true)]
    [string]$Title,

    [Parameter(Mandatory = $true)]
    [string]$Message,

    [Parameter(Mandatory = $true)]
    [string]$NotifyScriptPath
)

function Escape-TaskArgument {
    param([string]$Text)
    return $Text.Replace('`', '``').Replace('"', '`"')
}

$runAt = [datetime]::Parse($RunAtIso, [System.Globalization.CultureInfo]::InvariantCulture)
$escapedKind = Escape-TaskArgument $Kind
$escapedTitle = Escape-TaskArgument $Title
$escapedMessage = Escape-TaskArgument $Message
$escapedScript = Escape-TaskArgument $NotifyScriptPath

$arguments = "-NoProfile -WindowStyle Hidden -ExecutionPolicy Bypass -File `"$escapedScript`" -Kind `"$escapedKind`" -Title `"$escapedTitle`" -Message `"$escapedMessage`""
$action = New-ScheduledTaskAction -Execute 'powershell.exe' -Argument $arguments
$trigger = New-ScheduledTaskTrigger -Once -At $runAt
$principal = New-ScheduledTaskPrincipal -UserId $env:USERNAME -LogonType Interactive -RunLevel Limited
$settings = New-ScheduledTaskSettingsSet -AllowStartIfOnBatteries -DontStopIfGoingOnBatteries -StartWhenAvailable -MultipleInstances IgnoreNew

Register-ScheduledTask -TaskName $TaskName -Action $action -Trigger $trigger -Principal $principal -Settings $settings -Force | Out-Null
"""

    if not notify_script.exists() or notify_script.read_text(encoding='utf-8') != notify_script_content:
        notify_script.write_text(notify_script_content, encoding='utf-8')
    if not register_script.exists() or register_script.read_text(encoding='utf-8') != register_script_content:
        register_script.write_text(register_script_content, encoding='utf-8')

    return notify_script, register_script


def _register_windows_alert(kind, scheduled_time, message, title):
    if scheduled_time is None:
        return f"No {kind} time was provided."

    if scheduled_time.tzinfo is not None:
        scheduled_time = scheduled_time.astimezone().replace(tzinfo=None)

    now = datetime_module.datetime.now()
    if scheduled_time <= now:
        if kind == 'alarm':
            scheduled_time = scheduled_time + datetime_module.timedelta(days=1)
        else:
            return f"The requested {kind} time is in the past."

    notify_script, register_script = _ensure_system_alarm_scripts()
    task_name = f"Codex{kind.capitalize()}_{scheduled_time.strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}"
    run_at_iso = scheduled_time.strftime('%Y-%m-%dT%H:%M:%S')

    command = [
        'powershell.exe',
        '-NoProfile',
        '-ExecutionPolicy', 'Bypass',
        '-File', str(register_script),
        '-TaskName', task_name,
        '-RunAtIso', run_at_iso,
        '-Kind', kind,
        '-Title', title,
        '-Message', message,
        '-NotifyScriptPath', str(notify_script),
    ]

    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=30)
    except Exception as error:
        return f"Failed to register Windows {kind}: {error}"

    if result.returncode != 0:
        details = (result.stderr or result.stdout or '').strip()
        return f"Failed to register Windows {kind}: {details or 'unknown error'}"

    return f"{title} registered in Windows Task Scheduler for {scheduled_time.strftime('%I:%M %p').lstrip('0')}."


def _schedule_alarm(message, scheduled_time):
    label = str(message or 'Alarm').strip()
    if label.lower() in {'alarm', 'set alarm', 'set alarm for', 'alarm set'}:
        label = 'Alarm'
    return _register_windows_alert('alarm', scheduled_time, label, 'Codex Alarm')


def _schedule_reminder(message, delay_seconds):
    if delay_seconds is None:
        return 'No reminder duration was provided.'

    delay_seconds = max(0, float(delay_seconds))
    scheduled_time = datetime_module.datetime.now() + datetime_module.timedelta(seconds=delay_seconds)
    reminder_label = str(message or 'Reminder').strip()
    return _register_windows_alert('reminder', scheduled_time, reminder_label, 'Codex Reminder')


def _set_system_info(info_type):


    info_type = _normalize_text(info_type or "overview")

    if info_type in {"battery", "battery_percentage", "battery status"}:
        battery = psutil.sensors_battery()
        if battery is None:
            return "Battery information is not available on this device."
        plugged = "plugged in" if battery.power_plugged else "on battery"
        return f"Battery: {battery.percent}% ({plugged})"

    if info_type in {"internet", "connection", "network"}:
        try:
            socket.create_connection(("1.1.1.1", 53), timeout=2).close()
            return "Internet connection: available"
        except Exception:
            return "Internet connection: unavailable"

    if info_type in {"storage", "disk", "space", "drive"}:
        usage = shutil.disk_usage("C:\\")
        free_gb = usage.free / (1024**3)
        total_gb = usage.total / (1024**3)
        used_gb = usage.used / (1024**3)
        return f"C: drive - used {used_gb:.1f} GB / total {total_gb:.1f} GB, free {free_gb:.1f} GB"

    cpu = psutil.cpu_percent(interval=0.2)
    ram = psutil.virtual_memory().percent
    battery = psutil.sensors_battery()
    battery_text = f"{battery.percent}%" if battery else "N/A"
    return f"CPU {cpu}%, RAM {ram}%, Battery {battery_text}"


def _open_web_target(query=None, url=None, site="google"):
    if url:
        url_text = str(url).strip()
        if not re.match(r"^https?://", url_text, re.I):
            return "Only http and https URLs are allowed."
        webbrowser.open_new_tab(url_text)
        return f"Opened URL: {url_text}"

    query = str(query or "").strip()
    if not query:
        return "No search query was provided."

    site = _normalize_text(site)
    if site == "youtube":
        search_url = f"https://www.youtube.com/results?search_query={quote_plus(query)}"
    elif site == "bing":
        search_url = f"https://www.bing.com/search?q={quote_plus(query)}"
    else:
        search_url = f"https://www.google.com/search?q={quote_plus(query)}"

    webbrowser.open_new_tab(search_url)
    return f"Opened web search for: {query}"


def _handle_memory_box(intent_json, alias_map):
    alias_name = _get_parameter(intent_json, "alias_name") or _get_parameter(intent_json, "name")
    target_path = _get_parameter(intent_json, "target_path") or _get_parameter(intent_json, "path")
    if not alias_name or not target_path:
        return "Alias name or target path missing."

    alias_map[alias_name] = target_path
    try:
        with open(BASE_DIR / "shortcuts.json", "w", encoding="utf-8") as handle:
            json.dump(alias_map, handle, indent=4)
        return f"Saved shortcut '{alias_name}' -> '{target_path}'"
    except Exception as error:
        return f"Shortcut saved in memory but not on disk: {error}"


def _parse_human_datetime(text_value):
    normalized = _normalize_text(text_value)
    if not normalized:
        return None

    now = datetime_module.datetime.now()

    relative_match = re.search(r"\bin\s+(\d+)\s*(second|seconds|minute|minutes|hour|hours)\b", normalized)
    if relative_match:
        amount = int(relative_match.group(1))
        unit = relative_match.group(2)
        if unit.startswith("second"):
            return now + datetime_module.timedelta(seconds=amount)
        if unit.startswith("minute"):
            return now + datetime_module.timedelta(minutes=amount)
        if unit.startswith("hour"):
            return now + datetime_module.timedelta(hours=amount)

    time_match = re.search(r"\b(?:at\s*)?(\d{1,2})(?::(\d{2}))?\s*(am|pm)?\b", normalized)
    if time_match:
        hour = int(time_match.group(1))
        minute = int(time_match.group(2) or 0)
        meridiem = time_match.group(3)
        if meridiem == "pm" and hour != 12:
            hour += 12
        if meridiem == "am" and hour == 12:
            hour = 0

        day_offset = 1 if "tomorrow" in normalized else 0
        candidate = now.replace(hour=hour, minute=minute, second=0, microsecond=0)
        if day_offset:
            candidate += datetime_module.timedelta(days=1)
        elif candidate <= now:
            candidate += datetime_module.timedelta(days=1)
        return candidate

    if "tomorrow" in normalized:
        tomorrow = now + datetime_module.timedelta(days=1)
        return tomorrow.replace(hour=9, minute=0, second=0, microsecond=0)

    return None


def _extract_scheduled_datetime(intent_json):
    scheduled_time = _get_parameter(intent_json, "scheduled_time")
    if scheduled_time:
        if isinstance(scheduled_time, datetime_module.datetime):
            return scheduled_time
        text_value = str(scheduled_time).strip()
        parsed_natural = _parse_human_datetime(text_value)
        if parsed_natural is not None:
            return parsed_natural
        for pattern in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%H:%M",
            "%I:%M %p",
            "%I %p",
        ):
            try:
                parsed = datetime_module.datetime.strptime(text_value, pattern)
                if pattern in {"%H:%M", "%I:%M %p", "%I %p"}:
                    now = datetime_module.datetime.now()
                    parsed = now.replace(
                        hour=parsed.hour,
                        minute=parsed.minute,
                        second=0,
                        microsecond=0,
                    )
                    if parsed <= now:
                        parsed += datetime_module.timedelta(days=1)
                return parsed
            except Exception:
                continue

        try:
            parsed = datetime_module.datetime.fromisoformat(text_value)
            if parsed.tzinfo is not None:
                parsed = parsed.astimezone().replace(tzinfo=None)
            return parsed
        except Exception:
            return None

    duration_seconds = _get_parameter(intent_json, "duration_seconds")
    duration_minutes = _get_parameter(intent_json, "duration_minutes")
    if duration_seconds is not None:
        try:
            return datetime_module.datetime.now() + datetime_module.timedelta(seconds=float(duration_seconds))
        except Exception:
            return None
    if duration_minutes is not None:
        try:
            return datetime_module.datetime.now() + datetime_module.timedelta(minutes=float(duration_minutes))
        except Exception:
            return None
    return None


def _find_safest_base_folder(intent_json, alias_map):
    folder_value = _get_parameter(intent_json, "folder") or _get_parameter(intent_json, "path")
    if folder_value:
        alias_value = alias_map.get(str(folder_value))
        if alias_value:
            folder_value = alias_value
        resolved_folder = _safe_target_path(folder_value, default_base=DEFAULT_SAFE_DIR)
        if resolved_folder and resolved_folder.exists():
            return resolved_folder

    target_text = _normalize_text(_get_parameter(intent_json, "target"))
    if "project" in target_text:
        return BASE_DIR
    if "download" in target_text:
        return HOME_DIR / "Downloads"
    if "picture" in target_text or "image" in target_text:
        return HOME_DIR / "Pictures"
    if "desktop" in target_text:
        return HOME_DIR / "Desktop"
    return DEFAULT_SAFE_DIR


def verify_and_execute(intent_json, alias_map):
    mode = intent_json.get("mode", "action")
    intent = intent_json.get("intent", "unknown")
    confidence = intent_json.get("confidence", 0)
    action_type = _normalize_text(_get_parameter(intent_json, "action_type"))

    if intent == "file_action":
        if action_type in {"delete", "remove"}:
            intent = "delete_file"
        elif action_type == "move":
            intent = "move_file"
        elif action_type == "copy":
            intent = "copy_file"
        elif action_type in {"rename", "move_rename"}:
            intent = "rename_file"
        elif action_type in {"create", "new", "make"}:
            intent = "create_file"
        elif action_type in {"open", "open_file", "launch"}:
            intent = "open_file"
        elif action_type in {"open_folder", "folder"}:
            intent = "open_folder"
        elif action_type in {"open_parent_folder", "parent_folder", "up_folder"}:
            intent = "open_parent_folder"
        elif action_type in {"create_folder", "mkdir", "new_folder"}:
            intent = "create_folder"
        elif action_type in {"rename_folder", "folder_rename"}:
            intent = "rename_folder"
        elif action_type in {"move_folder", "folder_move"}:
            intent = "move_folder"
        elif action_type in {"copy_folder", "folder_copy"}:
            intent = "copy_folder"
        elif action_type in {"delete_folder", "remove_folder"}:
            intent = "delete_folder"
        elif action_type in {"list_folder", "show_folder", "contents"}:
            intent = "list_folder"
        elif action_type in {"search_files", "find_files", "search_folder"}:
            intent = "search_files"
        elif action_type == "find_latest":
            intent = "find_latest"
        elif action_type in {"upload", "upload_files", "read_upload"}:
            intent = "upload_files"

    if mode == "blocked":
        return intent_json.get("reply", "Request blocked.")

    if confidence < CONFIG["min_confidence_threshold"] and intent not in {"chat"}:
        return "Command rejected: insufficient model extraction confidence."

    text_fragments = [
        _get_parameter(intent_json, "target", ""),
        _get_parameter(intent_json, "path", ""),
        _get_parameter(intent_json, "folder", ""),
        _get_parameter(intent_json, "destination", ""),
        _get_parameter(intent_json, "url", ""),
        _get_parameter(intent_json, "command", ""),
    ]
    combined_text = " ".join(str(fragment) for fragment in text_fragments if fragment)
    normalized_combined = combined_text.lower()

    for blocked in CONFIG["blocked_paths"]:
        if blocked.lower() in normalized_combined or "system32" in normalized_combined:
            return f"Security intercept: unauthorized access attempt to protected path: {blocked}"
    if normalized_combined.endswith(".dll") or normalized_combined.endswith(".sys"):
        return "Security intercept: modifying system binaries is prohibited."

    if intent == "chat":
        reply_text = intent_json.get("reply", "Ready.")
        lower_reply = reply_text.lower()
        if any(keyword in lower_reply for keyword in ("time", "date", "clock")):
            now = datetime_module.datetime.now()
            if "date" in lower_reply and "time" not in lower_reply:
                return now.strftime("%A, %B %d, %Y")
            if "time" in lower_reply:
                return now.strftime("%H:%M:%S")
        return reply_text

    if intent == "memory_box":
        return _handle_memory_box(intent_json, alias_map)

    if intent == "open_app":
        target = _get_parameter(intent_json, "target") or _get_parameter(intent_json, "app_name")
        if not target:
            return "No application target was provided."
        resolved_target = _resolve_app_target(target)
        if _shell_open(resolved_target):
            return f"Opening application: {target}"
        return f"Could not open application: {target}"

    if intent == "close_app":
        target = _get_parameter(intent_json, "target") or _get_parameter(intent_json, "app_name")
        return _close_application(target)

    if intent == "open_folder":
        folder = _find_safest_base_folder(intent_json, alias_map)
        return _open_folder(folder)

    if intent == "open_parent_folder":
        folder_value = _resolve_alias_value(_get_parameter(intent_json, "path") or _get_parameter(intent_json, "folder"), alias_map)
        folder_path = _safe_target_path(folder_value) if folder_value else DEFAULT_SAFE_DIR
        if folder_path is None:
            return "No safe folder path was provided."
        parent_folder = Path(folder_path).parent
        return _open_folder(parent_folder)

    if intent == "open_file":
        file_value = _resolve_alias_value(_get_parameter(intent_json, "path") or _get_parameter(intent_json, "target"), alias_map)
        file_path = _safe_target_path(file_value)
        if file_path is None:
            return "No safe file path was provided."
        return _open_file(file_path)

    if intent == "create_folder":
        destination_value = _get_parameter(intent_json, "path") or _get_parameter(intent_json, "target_path")
        folder_name = _get_parameter(intent_json, "name") or _get_parameter(intent_json, "target")
        if destination_value:
            destination_value = _resolve_alias_value(destination_value, alias_map)
            destination = _safe_target_path(destination_value)
            if destination is None:
                return f"Blocked unsafe destination outside allowed folders: {destination_value}"
            if folder_name:
                destination = destination / Path(str(folder_name)).name
            return _create_folder(destination)

        folder_base = _find_safest_base_folder(intent_json, alias_map)
        if not folder_name:
            return "No folder name was provided."
        destination = folder_base / Path(str(folder_name)).name
        return _create_folder(destination)

    if intent == "create_file":
        file_name = _get_parameter(intent_json, "name") or _get_parameter(intent_json, "target")
        content = _get_parameter(intent_json, "content", "")
        folder = _find_safest_base_folder(intent_json, alias_map)
        if not file_name:
            return "No file name was provided."
        destination = folder / Path(str(file_name)).name
        return _create_file(destination, content)

    if intent == "rename_folder":
        source_value = _resolve_alias_value(_get_parameter(intent_json, "path") or _get_parameter(intent_json, "target"), alias_map)
        source_path = _safe_target_path(source_value)
        new_name = _get_parameter(intent_json, "name") or _get_parameter(intent_json, "destination")
        return _rename_folder(source_path, new_name)

    if intent in {"move_folder", "copy_folder"}:
        source_value = _resolve_alias_value(_get_parameter(intent_json, "path") or _get_parameter(intent_json, "target"), alias_map)
        source_path = _safe_target_path(source_value)
        destination_value = _resolve_alias_value(_get_parameter(intent_json, "destination") or _get_parameter(intent_json, "folder"), alias_map)
        destination_path = None
        if destination_value:
            destination_path = _safe_target_path(destination_value, default_base=DEFAULT_SAFE_DIR)
        operation = "move" if intent == "move_folder" else "copy"
        return _move_or_copy_folder(source_path, destination_path, operation)

    if intent == "delete_folder":
        folder_value = _resolve_alias_value(_get_parameter(intent_json, "path") or _get_parameter(intent_json, "target"), alias_map)
        folder_path = _safe_target_path(folder_value)
        if folder_path is None:
            return "No safe folder path was provided."
        return _delete_folder(folder_path)

    if intent == "list_folder":
        folder = _find_safest_base_folder(intent_json, alias_map)
        return _list_folder(folder)

    if intent == "search_files":
        folder = _find_safest_base_folder(intent_json, alias_map)
        query = _get_parameter(intent_json, "query") or _get_parameter(intent_json, "target")
        return _search_files(folder, query)

    if intent == "analyze_folder":
        folder = _resolve_analysis_folder(intent_json, alias_map)
        if folder is None:
            return "No safe study folder path was provided."
        request_text = _get_parameter(intent_json, "content") or _get_parameter(intent_json, "target") or "summarize"
        analysis_mode = _normalize_text(_get_parameter(intent_json, "analysis_mode") or request_text or "summary")
        return _generate_folder_study_response(folder, request_text, analysis_mode)

    if intent == "upload_files":
        upload_paths = _resolve_upload_paths(intent_json, alias_map)
        if not upload_paths:
            return "No safe upload file or folder path was provided."
        request_text = _get_parameter(intent_json, "content") or _get_parameter(intent_json, "query") or "summarize the content"
        return _generate_upload_response(upload_paths, request_text)

    if intent in {"move_file", "copy_file"}:
        source_value = _resolve_alias_value(_get_parameter(intent_json, "path") or _get_parameter(intent_json, "target"), alias_map)
        source_path = _safe_target_path(source_value)
        destination_value = _resolve_alias_value(_get_parameter(intent_json, "destination") or _get_parameter(intent_json, "folder"), alias_map)
        destination_path = None
        if destination_value:
            destination_path = _safe_target_path(destination_value, default_base=DEFAULT_SAFE_DIR)
        if destination_path and destination_path.is_dir():
            destination_path = destination_path / Path(str(source_path)).name
        operation = "move" if intent == "move_file" else "copy"
        return _move_or_copy_file(source_path, destination_path, operation)

    if intent == "delete_file":
        source_value = _resolve_alias_value(_get_parameter(intent_json, "path") or _get_parameter(intent_json, "target"), alias_map)
        source_path = _safe_target_path(source_value)
        if source_path is None:
            return "No safe file path was provided."
        if not source_path.exists():
            return f"File not found: {source_path}"
        try:
            send2trash(str(source_path))
            return f"Moved to Recycle Bin: {source_path}"
        except Exception as error:
            return f"Delete failed: {error}"

    if intent == "rename_file":
        source_value = _resolve_alias_value(_get_parameter(intent_json, "path") or _get_parameter(intent_json, "target"), alias_map)
        source_path = _safe_target_path(source_value)
        new_name = _get_parameter(intent_json, "name") or _get_parameter(intent_json, "destination")
        return _rename_file(source_path, new_name)

    if intent == "find_latest":
        folder = _find_safest_base_folder(intent_json, alias_map)
        target_text = _normalize_text(_get_parameter(intent_json, "target") or "")
        file_kind = None
        if any(keyword in target_text for keyword in ("video", "clip", "movie")):
            file_kind = "video"
        elif any(keyword in target_text for keyword in ("image", "picture", "photo", "screenshot")):
            file_kind = "image"
        return _find_latest_file(folder, file_kind=file_kind)

    if intent == "web_search":
        query = _get_parameter(intent_json, "query") or _get_parameter(intent_json, "target")
        site = _get_parameter(intent_json, "site", "google")
        return _open_web_target(query=query, site=site)

    if intent == "open_url":
        url = _get_parameter(intent_json, "url") or _get_parameter(intent_json, "target")
        if not url:
            return "No URL was provided."
        return _open_web_target(url=url)

    if intent == "note_add":
        note_text = _get_parameter(intent_json, "note_text") or _get_parameter(intent_json, "content") or _get_parameter(intent_json, "target")
        return _save_note(note_text)

    if intent == "note_read":
        return _read_note()

    if intent == "note_clear":
        return _clear_note()

    if intent in {"set_timer", "set_alarm", "set_reminder"}:
        message = _get_parameter(intent_json, "content") or _get_parameter(intent_json, "target") or intent.replace("_", " ")
        scheduled_time = _extract_scheduled_datetime(intent_json)
        if intent == "set_timer":
            if scheduled_time is None:
                return "No timer duration was provided."
            delay_seconds = (scheduled_time - datetime_module.datetime.now()).total_seconds()
            return _schedule_timer(message, delay_seconds)
        if intent == "set_alarm":
            return _schedule_alarm(message, scheduled_time)
        if scheduled_time is not None:
            delay_seconds = (scheduled_time - datetime_module.datetime.now()).total_seconds()
            return _schedule_reminder(message, delay_seconds)
        duration_seconds = _get_parameter(intent_json, "duration_seconds")
        duration_minutes = _get_parameter(intent_json, "duration_minutes")
        if duration_seconds is not None:
            return _schedule_reminder(message, float(duration_seconds))
        if duration_minutes is not None:
            return _schedule_reminder(message, float(duration_minutes) * 60)
        return "No reminder time was provided."

    if intent == "screenshot":
        scope = _get_parameter(intent_json, "screenshot_scope", "full")
        target = _get_parameter(intent_json, "target") or _get_parameter(intent_json, "window_title") or _get_parameter(intent_json, "app_name")
        return _capture_screen(
            scope="window" if _normalize_text(scope) == "window" else "full" if scope != "active" else "active",
            filename_prefix="screenshot",
            target=target,
            window_title=_get_parameter(intent_json, "window_title") or target,
        )

    if intent == "capture_active_window":
        return _capture_screen(scope="active", filename_prefix="active_window")

    if intent == "system_info":
        info_type = _get_parameter(intent_json, "info_type") or _get_parameter(intent_json, "target")
        return _set_system_info(info_type)

    return "Processed intent successfully but no executor match was found."
